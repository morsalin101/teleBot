from docx import Document
from docx.shared import Pt
import telebot
from fpdf import FPDF
import os

# Directly assign the bot token
BOT_TOKEN = '7181692028:AAHZg0n47h6hUs0VriLb24X_twNUZDBBXo'
bot = telebot.TeleBot(BOT_TOKEN)

# Define paths
TEMPLATE_PATH = r"C:\botFiles\cover.docx"
OUTPUT_PATH = r"C:\botFiles\filled_cover.docx"

# Define the placeholders that correspond to the data to be filled
placeholders = [
    "<Title>",
    "<Course Title>",
    "<T Name>",
    "<Designation>",  # Added Designation
    "<Name>",
    "<ID>",
    "<Section>",
    "<Semester>",
    "<Department>",
    "<date>"  # Added date question
]

@bot.message_handler(commands=['start'])
def send_welcome(message):
    # Send a small menu when the bot is started
    menu_text = "Welcome to the Document Bot! Please choose an option:\n\n"
    menu_text += "1. Convert .txt file to PDF (/TextFileToPdf)\n"
    menu_text += "2. To get Assignment Cover (/Acover)\n"
    menu_text += "Type the corresponding command or click on the command to proceed."
    bot.reply_to(message, menu_text)

# Command for PDF conversion
@bot.message_handler(commands=['TextFileToPdf'])
def handle_pdf_conversion(message):
    bot.reply_to(message, "Please upload a .txt file to convert it to PDF.")

# Command for filling out the assignment cover (replaces /form)
@bot.message_handler(commands=['Acover'])
def start_form(message):
    instructions = (
        "Please provide all the details in a single line, separated by commas:\n\n"
        "Title,CourseTitle,TeacherName,Designation,YourName,ID,Section,Semester,Department,SubmissionDate\n\n"
        "Example:\n"
        "AI101,IntroductionToAI,JohnDoe,Professor,AliceSmith,12345678,B,3rd,CS,2024-08-20\n\n"
        "Type /cancel to cancel the process."
    )
    bot.reply_to(message, instructions)

@bot.message_handler(commands=['cancel'])
def cancel_process(message):
    bot.reply_to(message, "Process cancelled.")

@bot.message_handler(content_types=['document'])
def handle_docs(message):
    document = message.document
    
    # If the file is a text file
    if document.mime_type == 'text/plain':
        file_info = bot.get_file(document.file_id)
        downloaded_file = bot.download_file(file_info.file_path)
        
        # Decode the content of the text file
        text_content = downloaded_file.decode('utf-8')

        # Create a PDF from the text content
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        for line in text_content.splitlines():
            pdf.cell(200, 10, txt=line, ln=True)

        # Get the original file name without the extension
        original_filename = os.path.splitext(document.file_name)[0]

        # Define the path to save the PDF on your local computer
        pdf_directory = r"C:\botPdf"
        if not os.path.exists(pdf_directory):
            os.makedirs(pdf_directory)  # Create the directory if it doesn't exist

        # Create the PDF file path with the same name as the original text file
        pdf_file_path = os.path.join(pdf_directory, f"{original_filename}.pdf")
        pdf.output(pdf_file_path)

        # Send the PDF file back to the user
        with open(pdf_file_path, 'rb') as pdf_file:
            bot.send_document(message.chat.id, pdf_file)
    
    else:
        bot.reply_to(message, "Please send a valid .txt file.")

@bot.message_handler(func=lambda message: True)
def handle_input(message):
    # Split the user input by commas
    user_input = [item.strip() for item in message.text.split(',')]

    if len(user_input) != len(placeholders):
        bot.reply_to(message, "Please provide all details in the correct format.")
        return

    data = dict(zip(placeholders, user_input))
    fill_template(data)

    # Send the filled .docx file
    with open(OUTPUT_PATH, 'rb') as filled_file:
        bot.send_document(message.chat.id, filled_file)

    bot.send_message(message.chat.id, "Your document has been filled and sent back.")

def fill_template(data):
    doc = Document(TEMPLATE_PATH)
    
    # Replace placeholders in the document with user-provided data and set font size to 18
    for para in doc.paragraphs:
        for key, value in data.items():
            if key in para.text:
                para.text = para.text.replace(key, value)
                
                # Set the font size to 18 for the entire paragraph
                for run in para.runs:
                    run.font.size = Pt(18)

    doc.save(OUTPUT_PATH)

# Start the bot's polling
bot.infinity_polling()
